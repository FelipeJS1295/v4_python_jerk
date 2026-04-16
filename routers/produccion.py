# app/routers/produccion.py
from fastapi import APIRouter, HTTPException, Request
from openpyxl.styles import Font, Alignment, PatternFill
from fastapi.responses import JSONResponse, HTMLResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
from typing import Optional, List, Dict, Any
from io import BytesIO
from datetime import datetime
import pandas as pd


from db import conectar_mysql

router = APIRouter(prefix="/produccion", tags=["Producción"])
templates = Jinja2Templates(directory="templates")


# =========================
# VISTAS
# =========================
@router.get("/", response_class=HTMLResponse)
async def vista_produccion(request: Request):
    return templates.TemplateResponse("produccion.html", {"request": request})


@router.get("/create", response_class=HTMLResponse)
async def vista_crear_produccion(request: Request):
    return templates.TemplateResponse("produccion/create.html", {"request": request})


@router.get("/create-reparacion", response_class=HTMLResponse)
async def vista_crear_reparacion(request: Request):
    return templates.TemplateResponse("produccion/create_reparacion.html", {"request": request})


@router.get("/edit/{id}", response_class=HTMLResponse)
async def vista_editar_produccion(request: Request, id: int):
    return templates.TemplateResponse("produccion/edit.html", {"request": request, "orden_id": id})


# =========================
# MODELOS
# =========================
class FiltroResumen(BaseModel):
    trabajador_id: int
    fecha_desde: Optional[str] = None
    fecha_hasta: Optional[str] = None


class OrdenProduccion(BaseModel):
    productos_id: int
    fecha: str
    numero_orden_trabajo: str


class ProduccionBatch(BaseModel):
    trabajador_id: int
    ordenes: List[OrdenProduccion]


class ProduccionCreate(BaseModel):
    productos_id: int
    numero_orden_trabajo: str
    fecha: str


class OrdenReparacionItem(BaseModel):
    producto_id: int
    fecha: str
    numero_orden_trabajo: str
    descripcion: Optional[str] = None
    precio_reparacion: Optional[float] = 0


class ReparacionCreate(BaseModel):
    trabajadores_id: int
    ordenes: List[OrdenReparacionItem]


# =========================
# HELPERS REUTILIZABLES
# =========================
def _obtener_resumen_trabajador(conn, filtro: FiltroResumen) -> Dict[str, Any]:
    """
    Devuelve el mismo payload que /resumen-detallado (JSON) para reutilizar en el endpoint de Excel.
    Estructura:
    {
      "trabajador": str,
      "rol": str,
      "fecha_desde": 'YYYY-MM-DD' | None,
      "fecha_hasta": 'YYYY-MM-DD' | None,
      "detalle": [ {fecha, numero_orden_trabajo, producto, tipo, descripcion, costo}, ... ]
    }
    """
    cursor = conn.cursor(dictionary=True)
    try:
        # Datos del trabajador y su rol
        cursor.execute("""
            SELECT u.rol, t.nombres 
            FROM users u 
            JOIN trabajadores t ON u.id = t.id 
            WHERE t.id = %s
        """, (filtro.trabajador_id,))
        user_data = cursor.fetchone()
        if not user_data:
            raise HTTPException(status_code=404, detail="Trabajador no encontrado")

        rol = (user_data["rol"] or "").strip()
        nombre = user_data["nombres"]

        # Mapeo de campo costo por rol
        campo_costo = {
            "Tapiceria": "costo_tapiceria",
            "Tapicería": "costo_tapiceria",
            "tapiceria": "costo_tapiceria",
            "TAPICERIA": "costo_tapiceria",
            "Costura": "costo_costura",
            "costura": "costo_costura",
            "COSTURA": "costo_costura",
            "Esqueleteria": "costo_esqueleteria",
            "esqueleteria": "costo_esqueleteria",
            "ESQUELETERIA": "costo_esqueleteria"
        }.get(rol, None)

        if not campo_costo:
            campo_costo = "precio_reparacion"

        # Verifica que la columna exista; si no, cae a precio_reparacion
        c = conn.cursor(dictionary=True)
        c.execute("SHOW COLUMNS FROM productos LIKE %s", (campo_costo,))
        columna_existe = c.fetchone()
        c.close()
        if not columna_existe and campo_costo != "precio_reparacion":
            campo_costo = "precio_reparacion"

        # Query para detalle
        if campo_costo == "precio_reparacion":
            # Si no existe costo por rol, usa siempre el precio de reparación
            query = """
                SELECT 
                    p.fecha,
                    p.numero_orden_trabajo,
                    pr.nombre AS producto,
                    p.tipo,
                    p.descripcion,
                    COALESCE(p.precio_reparacion, 0) AS costo
                FROM produccion p
                JOIN productos pr ON p.productos_id = pr.id
                WHERE p.trabajadores_id = %s
            """
        else:
            # Si hay costo por rol, aún así respeta reparaciones
            query = f"""
                SELECT 
                    p.fecha,
                    p.numero_orden_trabajo,
                    pr.nombre AS producto,
                    p.tipo,
                    p.descripcion,
                    CASE 
                    WHEN p.tipo = 'reparacion' THEN COALESCE(p.precio_reparacion, 0)
                    ELSE COALESCE(pr.{campo_costo}, 0)
                    END AS costo
                FROM produccion p
                JOIN productos pr ON p.productos_id = pr.id
                WHERE p.trabajadores_id = %s
            """

        params: List[Any] = [filtro.trabajador_id]
        if filtro.fecha_desde:
            query += " AND p.fecha >= %s"
            params.append(filtro.fecha_desde)
        if filtro.fecha_hasta:
            query += " AND p.fecha <= %s"
            params.append(filtro.fecha_hasta)
        query += " ORDER BY p.fecha DESC"

        cursor.execute(query, params)
        datos = cursor.fetchall()

        for d in datos:
            if d["fecha"]:
                d["fecha"] = d["fecha"].strftime("%Y-%m-%d")

        return {
            "trabajador": nombre,
            "rol": rol,
            "fecha_desde": filtro.fecha_desde,
            "fecha_hasta": filtro.fecha_hasta,
            "detalle": datos
        }

    finally:
        cursor.close()


# =========================
# ENDPOINTS JSON
# =========================
@router.get("/ordenes", response_class=JSONResponse)
def obtener_ordenes_produccion():
    conn = conectar_mysql()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("""
            SELECT 
                p.id,
                p.numero_orden_trabajo,
                pr.nombre AS producto,
                t.nombres AS trabajador,
                p.fecha,
                p.tipo
            FROM produccion p
            JOIN productos pr ON p.productos_id = pr.id
            JOIN trabajadores t ON p.trabajadores_id = t.id
            ORDER BY p.fecha DESC
        """)
        resultados = cursor.fetchall()
        for r in resultados:
            if r["fecha"]:
                r["fecha"] = r["fecha"].strftime("%Y-%m-%d")
            if r["tipo"] == "normal":
                r["estado"] = "Pendiente"
            elif r["tipo"] == "reparacion":
                r["estado"] = "Reparado"
            else:
                r["estado"] = "Desconocido"
        return resultados
    finally:
        cursor.close()
        conn.close()


@router.post("/resumen-trabajador", response_class=JSONResponse)
def resumen_por_trabajador(filtro: FiltroResumen):
    conn = conectar_mysql()
    cursor = conn.cursor(dictionary=True)
    try:
        query = """
            SELECT 
                p.numero_orden_trabajo,
                pr.nombre AS producto,
                p.fecha,
                p.tipo,
                p.precio_reparacion,
                t.nombres AS trabajador
            FROM produccion p
            JOIN productos pr ON p.productos_id = pr.id
            JOIN trabajadores t ON p.trabajadores_id = t.id
            WHERE p.trabajadores_id = %s
        """
        params: List[Any] = [filtro.trabajador_id]

        if filtro.fecha_desde:
            query += " AND p.fecha >= %s"
            params.append(filtro.fecha_desde)
        if filtro.fecha_hasta:
            query += " AND p.fecha <= %s"
            params.append(filtro.fecha_hasta)

        query += " ORDER BY p.fecha DESC"

        cursor.execute(query, params)
        resultados = cursor.fetchall()

        total_unidades = len(resultados)
        total_monto = sum(r["precio_reparacion"] or 0 for r in resultados)

        for r in resultados:
            if r["fecha"]:
                r["fecha"] = r["fecha"].strftime("%Y-%m-%d")

        return {
            "detalle": resultados,
            "total_unidades": total_unidades,
            "total_monto": total_monto
        }
    finally:
        cursor.close()
        conn.close()


@router.post("/resumen-detallado", response_class=JSONResponse)
def resumen_detallado_trabajador(filtro: FiltroResumen):
    conn = conectar_mysql()
    try:
        data = _obtener_resumen_trabajador(conn, filtro)
        return data
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error interno: {str(e)}")
    finally:
        conn.close()


@router.get("/trabajadores", response_class=JSONResponse)
def obtener_trabajadores():
    conn = conectar_mysql()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("""
            SELECT t.id, t.nombres
            FROM trabajadores t
            JOIN users u ON t.id = u.id
            WHERE u.rol IN ('Tapiceria', 'Costura', 'Esqueleteria')
        """)
        trabajadores = cursor.fetchall()
        return trabajadores
    finally:
        cursor.close()
        conn.close()


@router.post("/crear", response_class=JSONResponse)
def crear_produccion_batch(data: ProduccionBatch):
    conn = conectar_mysql()
    cursor = conn.cursor(dictionary=True)
    try:
        # Rol del trabajador
        cursor.execute("""
            SELECT u.rol 
            FROM users u 
            JOIN trabajadores t ON u.id = t.id 
            WHERE t.id = %s
        """, (data.trabajador_id,))
        trabajador_info = cursor.fetchone()
        if not trabajador_info:
            raise HTTPException(status_code=404, detail="Trabajador no encontrado")

        rol_trabajador = trabajador_info["rol"]

        # Validaciones de duplicados
        numeros_orden = [orden.numero_orden_trabajo for orden in data.ordenes]
        if len(numeros_orden) != len(set(numeros_orden)):
            raise HTTPException(status_code=400, detail="No puede repetir números de orden en la misma solicitud")

        placeholders = ",".join(["%s"] * len(numeros_orden))
        cursor.execute(f"""
            SELECT p.numero_orden_trabajo
            FROM produccion p
            JOIN trabajadores t ON p.trabajadores_id = t.id
            JOIN users u ON t.id = u.id
            WHERE u.rol = %s AND p.numero_orden_trabajo IN ({placeholders})
        """, [rol_trabajador] + numeros_orden)
        ordenes_existentes = cursor.fetchall()
        if ordenes_existentes:
            numeros_duplicados = [o["numero_orden_trabajo"] for o in ordenes_existentes]
            raise HTTPException(
                status_code=400,
                detail=f"Los siguientes números de orden ya existen para trabajadores de {rol_trabajador}: {', '.join(numeros_duplicados)}"
            )

        # Inserción
        for orden in data.ordenes:
            cursor.execute("""
                INSERT INTO produccion (
                    trabajadores_id, productos_id, fecha, numero_orden_trabajo, tipo, created_at, updated_at
                ) VALUES (%s, %s, %s, %s, 'normal', NOW(), NOW())
            """, (
                data.trabajador_id,
                orden.productos_id,
                orden.fecha,
                orden.numero_orden_trabajo
            ))

        conn.commit()
        return {"mensaje": "Producción registrada correctamente"}
    except HTTPException:
        conn.rollback()
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=f"Error al registrar producción: {str(e)}")
    finally:
        cursor.close()
        conn.close()


@router.get("/orden/{id}", response_class=JSONResponse)
def obtener_orden(id: int):
    conn = conectar_mysql()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("""
            SELECT p.*, t.nombres AS trabajador_nombre, pr.nombre AS producto_nombre
            FROM produccion p
            JOIN trabajadores t ON p.trabajadores_id = t.id
            JOIN productos pr ON p.productos_id = pr.id
            WHERE p.id = %s
        """, (id,))
        data = cursor.fetchone()
        if not data:
            raise HTTPException(status_code=404, detail="Orden no encontrada")

        if data.get("fecha"):
            data["fecha"] = data["fecha"].strftime("%Y-%m-%d")

        return data
    finally:
        cursor.close()
        conn.close()


@router.put("/{id}", response_class=JSONResponse)
def actualizar_orden(id: int, data: ProduccionCreate):
    conn = conectar_mysql()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("""
            SELECT p.numero_orden_trabajo, p.trabajadores_id, u.rol
            FROM produccion p
            JOIN trabajadores t ON p.trabajadores_id = t.id
            JOIN users u ON t.id = u.id
            WHERE p.id = %s
        """, (id,))
        orden_actual = cursor.fetchone()
        if not orden_actual:
            raise HTTPException(status_code=404, detail="Orden no encontrada")

        # Validar cambio de número
        if orden_actual["numero_orden_trabajo"] != data.numero_orden_trabajo:
            rol_trabajador = orden_actual["rol"]
            cursor.execute("""
                SELECT p.id
                FROM produccion p
                JOIN trabajadores t ON p.trabajadores_id = t.id
                JOIN users u ON t.id = u.id
                WHERE u.rol = %s AND p.numero_orden_trabajo = %s AND p.id != %s
            """, (rol_trabajador, data.numero_orden_trabajo, id))
            orden_existente = cursor.fetchone()
            if orden_existente:
                raise HTTPException(
                    status_code=400,
                    detail=f"El número de orden {data.numero_orden_trabajo} ya existe para otro trabajador de {rol_trabajador}"
                )

        cursor.execute("""
            UPDATE produccion
            SET productos_id = %s, numero_orden_trabajo = %s, fecha = %s, updated_at = NOW()
            WHERE id = %s
        """, (
            data.productos_id,
            data.numero_orden_trabajo,
            data.fecha,
            id
        ))
        conn.commit()
        return {"mensaje": "Producción actualizada correctamente"}
    except HTTPException:
        conn.rollback()
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()
        conn.close()


@router.post("/crear-reparacion", response_class=JSONResponse)
def crear_reparaciones(data: ReparacionCreate):
    conn = conectar_mysql()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("""
            SELECT u.rol 
            FROM users u 
            JOIN trabajadores t ON u.id = t.id 
            WHERE t.id = %s
        """, (data.trabajadores_id,))
        trabajador_info = cursor.fetchone()
        if not trabajador_info:
            raise HTTPException(status_code=404, detail="Trabajador no encontrado")

        rol_trabajador = trabajador_info["rol"]

        numeros_orden = [orden.numero_orden_trabajo for orden in data.ordenes]
        if len(numeros_orden) != len(set(numeros_orden)):
            raise HTTPException(status_code=400, detail="No puede repetir números de orden en la misma solicitud")

        placeholders = ",".join(["%s"] * len(numeros_orden))
        cursor.execute(f"""
            SELECT p.numero_orden_trabajo
            FROM produccion p
            JOIN trabajadores t ON p.trabajadores_id = t.id
            JOIN users u ON t.id = u.id
            WHERE u.rol = %s AND p.numero_orden_trabajo IN ({placeholders})
        """, [rol_trabajador] + numeros_orden)
        ordenes_existentes = cursor.fetchall()
        if ordenes_existentes:
            numeros_duplicados = [o["numero_orden_trabajo"] for o in ordenes_existentes]
            raise HTTPException(
                status_code=400,
                detail=f"Los siguientes números de orden ya existen para trabajadores de {rol_trabajador}: {', '.join(numeros_duplicados)}"
            )

        for orden in data.ordenes:
            cursor.execute("""
                INSERT INTO produccion (
                    trabajadores_id, productos_id, fecha, numero_orden_trabajo, descripcion, precio_reparacion,
                    tipo, created_at, updated_at
                ) VALUES (%s, %s, %s, %s, %s, %s, 'reparacion', NOW(), NOW())
            """, (
                data.trabajadores_id,
                orden.producto_id,
                orden.fecha,
                orden.numero_orden_trabajo,
                orden.descripcion,
                orden.precio_reparacion
            ))

        conn.commit()
        return {"mensaje": "Reparaciones registradas correctamente"}
    except HTTPException:
        conn.rollback()
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=f"Error al guardar reparaciones: {str(e)}")
    finally:
        cursor.close()
        conn.close()


@router.get("/productos/", response_class=JSONResponse)
async def obtener_productos():
    conn = conectar_mysql()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("SELECT id, nombre, sku FROM productos ORDER BY nombre")
        productos = cursor.fetchall()
        return productos
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()
        conn.close()


@router.delete("/eliminar/{id}", response_class=JSONResponse)
def eliminar_orden(id: int):
    conn = conectar_mysql()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM produccion WHERE id = %s", (id,))
        conn.commit()
        return {"mensaje": "Orden eliminada correctamente"}
    finally:
        cursor.close()
        conn.close()


@router.post("/validar-numero-orden", response_class=JSONResponse)
def validar_numero_orden(data: dict):
    """Validar si un número de orden ya existe para un trabajador del mismo rol"""
    conn = conectar_mysql()
    cursor = conn.cursor(dictionary=True)
    try:
        trabajador_id = data.get("trabajador_id")
        numero_orden = data.get("numero_orden_trabajo")
        orden_id = data.get("orden_id", None)

        if not trabajador_id or not numero_orden:
            return {"valido": False, "mensaje": "Datos incompletos"}

        cursor.execute("""
            SELECT u.rol 
            FROM users u 
            JOIN trabajadores t ON u.id = t.id 
            WHERE t.id = %s
        """, (trabajador_id,))
        trabajador_info = cursor.fetchone()
        if not trabajador_info:
            return {"valido": False, "mensaje": "Trabajador no encontrado"}

        rol_trabajador = trabajador_info["rol"]

        if orden_id:
            cursor.execute("""
                SELECT p.numero_orden_trabajo, t.nombres
                FROM produccion p
                JOIN trabajadores t ON p.trabajadores_id = t.id
                JOIN users u ON t.id = u.id
                WHERE u.rol = %s AND p.numero_orden_trabajo = %s AND p.id != %s
                LIMIT 1
            """, (rol_trabajador, numero_orden, orden_id))
        else:
            cursor.execute("""
                SELECT p.numero_orden_trabajo, t.nombres
                FROM produccion p
                JOIN trabajadores t ON p.trabajadores_id = t.id
                JOIN users u ON t.id = u.id
                WHERE u.rol = %s AND p.numero_orden_trabajo = %s
                LIMIT 1
            """, (rol_trabajador, numero_orden))

        orden_existente = cursor.fetchone()
        if orden_existente:
            return {
                "valido": False,
                "mensaje": f"El número de orden '{numero_orden}' ya está asignado a {orden_existente['nombres']} ({rol_trabajador})"
            }

        return {"valido": True, "mensaje": "Número de orden disponible"}
    except Exception as e:
        return {"valido": False, "mensaje": f"Error al validar: {str(e)}"}
    finally:
        cursor.close()
        conn.close()


@router.get("/precios/{tipo}", response_class=JSONResponse)
def obtener_precios(tipo: str):
    tipo_campo = {
        "costura": "costo_costura",
        "tapiceria": "costo_tapiceria",
        "esqueleteria": "costo_esqueleteria"
    }
    if tipo not in tipo_campo:
        return JSONResponse(content=[], status_code=400)

    campo = tipo_campo[tipo]
    conn = conectar_mysql()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(f"""
            SELECT nombre, {campo} as costo
            FROM productos
            WHERE {campo} IS NOT NULL AND {campo} > 0
            ORDER BY nombre
        """)
        resultados = cursor.fetchall()
        return resultados
    finally:
        cursor.close()
        conn.close()


# =========================
# ENDPOINT XLSX (NUEVO)
# =========================
@router.post("/resumen-detallado/excel")
def resumen_detallado_excel(filtro: FiltroResumen):
    conn = conectar_mysql()
    try:
        data = _obtener_resumen_trabajador(conn, filtro)
        detalle = data.get("detalle", [])
        if not detalle:
            raise HTTPException(status_code=404, detail="No hay datos para exportar en el rango seleccionado.")

        # Detalle -> DataFrame
        df = pd.DataFrame(detalle)
        cols = ["fecha", "numero_orden_trabajo", "producto", "tipo", "descripcion", "costo"]
        for c in cols:
            if c not in df.columns:
                df[c] = None
        df = df[cols]

        # Normalización de valores
        df["fecha"] = pd.to_datetime(df["fecha"], errors="coerce").dt.strftime("%Y-%m-%d")
        df["tipo"] = df["tipo"].map(lambda x: "Normal" if str(x).lower() == "normal" else "Reparación")
        df["costo"] = pd.to_numeric(df["costo"], errors="coerce").fillna(0).astype(float)

        total = float(df["costo"].sum())
        promedio = float(df["costo"].mean()) if len(df) else 0.0

        # Hoja Resumen
        resumen_rows = [
            {"Campo": "Trabajador", "Valor": data.get("trabajador", "")},
            {"Campo": "Rol", "Valor": data.get("rol", "")},
            {"Campo": "Período", "Valor": f"{data.get('fecha_desde','')} a {data.get('fecha_hasta','')}"},
            {"Campo": "Total Órdenes", "Valor": len(df)},
            {"Campo": "Total Costos", "Valor": total},
            {"Campo": "Promedio por Orden", "Valor": promedio},
        ]
        df_resumen = pd.DataFrame(resumen_rows)

        # Excel en memoria
        output = BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Detalle")
            df_resumen.to_excel(writer, index=False, sheet_name="Resumen")

            # Ajustar anchos de columnas
            for ws in writer.sheets.values():
                for col_cells in ws.columns:
                    max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
                    ws.column_dimensions[col_cells[0].column_letter].width = min(max(10, max_len + 2), 60)

        output.seek(0)

        trabajador_safe = (data.get("trabajador") or "trabajador").replace(" ", "_")
        desde = (data.get("fecha_desde") or "").replace("-", "")
        hasta = (data.get("fecha_hasta") or "").replace("-", "")
        filename = f"resumen_{trabajador_safe}_{desde}_{hasta}.xlsx" if (desde or hasta) else f"resumen_{trabajador_safe}.xlsx"

        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers=headers
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al generar Excel: {str(e)}")
    finally:
        conn.close()

@router.post("/exportar-excel")
def exportar_produccion_excel(filtro: FiltroResumen):
    conn = conectar_mysql()
    try:
        # Reutilizamos la lógica de obtención de datos detallados
        data = _obtener_resumen_trabajador(conn, filtro)
        detalle = data.get("detalle", [])
        
        if not detalle:
            raise HTTPException(status_code=404, detail="No se encontraron registros para exportar.")

        # Crear DataFrame
        df = pd.DataFrame(detalle)
        
        # Renombrar columnas para el usuario final
        columnas_map = {
            "fecha": "Fecha",
            "numero_orden_trabajo": "N° Orden",
            "producto": "Producto / Modelo",
            "tipo": "Tipo",
            "descripcion": "Descripción/Observación",
            "costo": "Costo Unitario ($)"
        }
        df = df.rename(columns=columnas_map)
        
        # Asegurar que el costo sea numérico
        df["Costo Unitario ($)"] = pd.to_numeric(df["Costo Unitario ($)"], errors="coerce").fillna(0)

        # Crear el archivo Excel en memoria
        output = BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Detalle de Producción")
            
            # --- Formato Estético ---
            workbook = writer.book
            worksheet = writer.sheets["Detalle de Producción"]
            
            # Estilo para el encabezado
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            # Ajustar ancho de columnas automáticamente
            for col in worksheet.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except: pass
                worksheet.column_dimensions[column].width = max_length + 5

            # Hoja de Resumen Totales
            resumen_data = [
                ["Reporte de Producción - Jamaroff"],
                ["Trabajador:", data.get("trabajador")],
                ["Rol:", data.get("rol")],
                ["Desde:", filtro.fecha_desde or "Inicio"],
                ["Hasta:", filtro.fecha_hasta or "Hoy"],
                [""],
                ["TOTAL ÓRDENES:", len(df)],
                ["TOTAL A PAGAR:", df["Costo Unitario ($)"].sum()]
            ]
            df_resumen = pd.DataFrame(resumen_data)
            df_resumen.to_excel(writer, index=False, header=False, sheet_name="Resumen")

        output.seek(0)
        
        # Nombre del archivo dinámico
        nombre_archivo = f"Produccion_{data.get('trabajador').replace(' ', '_')}_{datetime.now().strftime('%d-%m-%Y')}.xlsx"
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
        )

    except Exception as e:
        print(f"Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()